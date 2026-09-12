from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import distinct, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import EngagementEvent, EngagementKind, User
from app.services.users import display_start_source, get_user


@dataclass(frozen=True, slots=True)
class EngagementCounts:
    recipes: int = 0
    lessons: int = 0


@dataclass(frozen=True, slots=True)
class UserRecord:
    user: User
    engagement: EngagementCounts


@dataclass(frozen=True, slots=True)
class UserPage:
    items: tuple[UserRecord, ...]
    page: int
    pages: int
    total: int


@dataclass(frozen=True, slots=True)
class StatisticsOverview:
    total_users: int
    phone_verified: int
    subscription_verified: int
    reachable: int
    unreachable: int
    new_24h: int
    new_7d: int
    new_30d: int
    languages: dict[str, int]
    sources: tuple[tuple[str, int], ...]
    recipe_views: int
    recipe_viewers: int
    lesson_views: int
    lesson_viewers: int
    top_recipes: tuple[tuple[int, int], ...]
    top_lessons: tuple[tuple[int, int], ...]


async def record_engagement(
    session: AsyncSession,
    *,
    telegram_id: int,
    kind: EngagementKind,
    target_id: int,
) -> None:
    user = await get_user(session, telegram_id)
    if user is None:
        return
    session.add(EngagementEvent(user_id=user.id, kind=kind, target_id=target_id))


async def _event_totals(session: AsyncSession, kind: EngagementKind) -> tuple[int, int]:
    row = (
        await session.execute(
            select(
                func.count(EngagementEvent.id),
                func.count(distinct(EngagementEvent.user_id)),
            ).where(EngagementEvent.kind == kind)
        )
    ).one()
    return int(row[0] or 0), int(row[1] or 0)


async def _top_content(
    session: AsyncSession, kind: EngagementKind, limit: int = 5
) -> tuple[tuple[int, int], ...]:
    rows = await session.execute(
        select(EngagementEvent.target_id, func.count(EngagementEvent.id).label("views"))
        .where(EngagementEvent.kind == kind)
        .group_by(EngagementEvent.target_id)
        .order_by(func.count(EngagementEvent.id).desc(), EngagementEvent.target_id.asc())
        .limit(limit)
    )
    return tuple((int(row[0]), int(row[1])) for row in rows)


async def get_statistics_overview(session: AsyncSession) -> StatisticsOverview:
    now = datetime.now(UTC)
    first_seen = func.coalesce(User.first_started_at, User.created_at)

    async def count_where(*conditions: object) -> int:
        return int(await session.scalar(select(func.count(User.id)).where(*conditions)) or 0)

    total = await count_where()
    phone_verified = await count_where(User.phone_verified_at.is_not(None))
    subscription_verified = await count_where(User.subscription_verified_at.is_not(None))
    reachable = await count_where(User.is_reachable.is_(True))
    language_rows = await session.execute(
        select(User.language_code, func.count(User.id)).group_by(User.language_code)
    )
    languages = {
        (row[0].value if row[0] is not None else "unknown"): int(row[1]) for row in language_rows
    }
    source_rows = await session.execute(
        select(User.first_start_source, func.count(User.id))
        .group_by(User.first_start_source)
        .order_by(func.count(User.id).desc(), User.first_start_source.asc())
        .limit(10)
    )
    combined_sources: dict[str, int] = {}
    for source, count in source_rows:
        label = display_start_source(source)
        combined_sources[label] = combined_sources.get(label, 0) + int(count)
    sources = tuple(sorted(combined_sources.items(), key=lambda item: (-item[1], item[0])))
    recipe_views, recipe_viewers = await _event_totals(session, EngagementKind.RECIPE)
    lesson_views, lesson_viewers = await _event_totals(session, EngagementKind.LESSON)
    return StatisticsOverview(
        total_users=total,
        phone_verified=phone_verified,
        subscription_verified=subscription_verified,
        reachable=reachable,
        unreachable=total - reachable,
        new_24h=await count_where(first_seen >= now - timedelta(hours=24)),
        new_7d=await count_where(first_seen >= now - timedelta(days=7)),
        new_30d=await count_where(first_seen >= now - timedelta(days=30)),
        languages=languages,
        sources=sources,
        recipe_views=recipe_views,
        recipe_viewers=recipe_viewers,
        lesson_views=lesson_views,
        lesson_viewers=lesson_viewers,
        top_recipes=await _top_content(session, EngagementKind.RECIPE),
        top_lessons=await _top_content(session, EngagementKind.LESSON),
    )


async def _engagement_for_users(
    session: AsyncSession, user_ids: tuple[int, ...]
) -> dict[int, EngagementCounts]:
    if not user_ids:
        return {}
    rows = await session.execute(
        select(EngagementEvent.user_id, EngagementEvent.kind, func.count(EngagementEvent.id))
        .where(EngagementEvent.user_id.in_(user_ids))
        .group_by(EngagementEvent.user_id, EngagementEvent.kind)
    )
    mutable: dict[int, dict[EngagementKind, int]] = {}
    for user_id, kind, count in rows:
        mutable.setdefault(int(user_id), {})[kind] = int(count)
    return {
        user_id: EngagementCounts(
            recipes=counts.get(EngagementKind.RECIPE, 0),
            lessons=counts.get(EngagementKind.LESSON, 0),
        )
        for user_id, counts in mutable.items()
    }


async def list_user_records(session: AsyncSession, *, page: int, page_size: int = 8) -> UserPage:
    total = int(await session.scalar(select(func.count(User.id))) or 0)
    pages = max(1, math.ceil(total / page_size))
    normalized_page = min(max(page, 0), pages - 1)
    users = tuple(
        (
            await session.scalars(
                select(User)
                .order_by(
                    func.coalesce(User.first_started_at, User.created_at).desc(),
                    User.id.desc(),
                )
                .offset(normalized_page * page_size)
                .limit(page_size)
            )
        ).all()
    )
    counts = await _engagement_for_users(session, tuple(user.id for user in users))
    return UserPage(
        items=tuple(
            UserRecord(user=user, engagement=counts.get(user.id, EngagementCounts()))
            for user in users
        ),
        page=normalized_page,
        pages=pages,
        total=total,
    )


async def get_user_record(session: AsyncSession, user_id: int) -> UserRecord | None:
    user = await session.get(User, user_id)
    if user is None:
        return None
    counts = await _engagement_for_users(session, (user.id,))
    return UserRecord(user=user, engagement=counts.get(user.id, EngagementCounts()))


async def all_user_records(session: AsyncSession) -> tuple[UserRecord, ...]:
    users = tuple((await session.scalars(select(User).order_by(User.id.asc()))).all())
    counts = await _engagement_for_users(session, tuple(user.id for user in users))
    return tuple(
        UserRecord(user=user, engagement=counts.get(user.id, EngagementCounts())) for user in users
    )


EXPORT_HEADERS = (
    "Telegram ID",
    "Username",
    "First name",
    "Last name",
    "Full name",
    "Phone",
    "Language",
    "First started (UTC)",
    "Last started (UTC)",
    "First start source",
    "Last start source",
    "Phone verified (UTC)",
    "Subscription last verified (UTC)",
    "Reachable",
    "Recipe deliveries",
    "Lesson views",
)


def _naive_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    return value


def _record_values(record: UserRecord) -> tuple[object, ...]:
    user = record.user
    full_name = " ".join(part for part in (user.first_name, user.last_name) if part)
    return (
        str(user.telegram_id),
        user.username or "",
        user.first_name or "",
        user.last_name or "",
        full_name,
        user.phone_number or "",
        user.language_code.value if user.language_code is not None else "",
        _naive_utc(user.first_started_at or user.created_at),
        _naive_utc(user.last_started_at),
        display_start_source(user.first_start_source),
        display_start_source(user.last_start_source),
        _naive_utc(user.phone_verified_at),
        _naive_utc(user.subscription_verified_at),
        "Yes" if user.is_reachable else "No",
        record.engagement.recipes,
        record.engagement.lessons,
    )


def _csv_safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return value


def build_users_csv(records: tuple[UserRecord, ...]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(EXPORT_HEADERS)
    for record in records:
        writer.writerow(tuple(_csv_safe(value) for value in _record_values(record)))
    return b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")


def build_users_xlsx(records: tuple[UserRecord, ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(EXPORT_HEADERS)
    for record in records:
        sheet.append(_record_values(record))
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for index in (8, 9, 12, 13):
            row[index - 1].number_format = "yyyy-mm-dd hh:mm:ss"
        for index in (1, 2, 3, 4, 5, 6, 10, 11):
            row[index - 1].number_format = "@"
    sheet.auto_filter.ref = f"A1:P{max(sheet.max_row, 1)}"
    for column_index, header in enumerate(EXPORT_HEADERS, start=1):
        values = [
            str(sheet.cell(row=row, column=column_index).value or "")
            for row in range(1, sheet.max_row + 1)
        ]
        width = min(max(len(header), *(len(value) for value in values)) + 2, 40)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
